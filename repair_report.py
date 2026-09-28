from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\Jaslyn\Desktop\fyp2_ - Copy.docx")
OUTPUT = Path(r"C:\Users\Jaslyn\Desktop\final_fyp_midi_gpt\fyp2_repaired.docx")


def insert_paragraph_before(paragraph, text, style):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = style
    new_paragraph.add_run(text)
    return new_paragraph


doc = Document(SOURCE)

# Keep the body and references intact while fixing confirmed mismatches.
for paragraph in doc.paragraphs:
    text = paragraph.text
    if "Table 6.2.1 shows the summary of the MIDI-GPT" in text:
        for run in paragraph.runs:
            run.text = run.text.replace("Table 6.2.1", "Table 6.5")
    if text.strip() == "Figure 5.18 Saving Generated Music History":
        for run in paragraph.runs:
            run.text = run.text.replace("Saving Generated Music History", "Saving Chatbot History")
    if "customize" in text:
        for run in paragraph.runs:
            run.text = run.text.replace("customize", "customise")
    if "MIDI-GPT is reviewed as a related transformer-based" in text:
        for run in paragraph.runs:
            run.text = run.text.replace("[17]", "[16]")
    if "DeBERTa had performed a high accuracy" in text:
        for run in paragraph.runs:
            run.text = run.text.replace("[31]", "[30]")
    if text.startswith("\u200c[21]"):
        for run in paragraph.runs:
            run.text = run.text.lstrip("\u200c")

# Some caption references are split across multiple Word runs.
for paragraph in doc.paragraphs:
    if "Table 6.2.1 shows the summary of the MIDI-GPT" in paragraph.text:
        full_text = paragraph.text.replace("Table 6.2.1", "Table 6.5")
        paragraph.runs[0].text = full_text
        for run in paragraph.runs[1:]:
            run.text = ""

# The TOC expects this subsection, and Figure 3.3 is its content.
if not any(p.text.strip() == "3.2.2 Use Case Diagram" for p in doc.paragraphs):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("Figure 3.3 Use Case Diagram"):
            insert_paragraph_before(paragraph, "3.2.2 Use Case Diagram", "Heading 3")
            break

# Ask Word to refresh TOC, captions and page fields when the file opens.
settings = doc.settings.element
update_fields = settings.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields")
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")

doc.save(OUTPUT)
print(OUTPUT)
